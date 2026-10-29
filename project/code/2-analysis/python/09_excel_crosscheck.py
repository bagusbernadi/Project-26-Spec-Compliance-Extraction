"""Langkah 7b: cek silang dengan Excel. Hasil ukur bersih + dim_spec ditulis ke crosscheck.xlsx; putusan lolos/gagal per baris dan pass rate
per parameter dan model dihitung dengan RUMUS Excel (MAXIFS, SUMIFS, COUNTIFS), dihitung ulang oleh LibreOffice, lalu dibandingkan dengan hasil SQL.
Batas kosong ditulis -1E+9 / 1E+9 di lembar spec agar rumus Excel sederhana. Semua baris harus PASS."""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
import common as C  # noqa: E402

m = pd.read_csv(C.CLEAN_DIR / "measurement_clean.csv", parse_dates=["test_date"])
s = pd.read_csv(C.RES_PREP / "dim_spec.csv", parse_dates=["valid_from"])
ref = pd.read_csv(C.RES_ANALYSIS / "passrate_param_model.csv")

wb = Workbook()
ws = wb.active; ws.title = "spec"
ws.append(["model", "parameter", "valid_from", "revision", "lower", "upper"])
for r in s.itertuples():
    ws.append([r.model, r.parameter, r.valid_from.to_pydatetime(), r.revision, -1e9 if pd.isna(r.lower_limit) else r.lower_limit, 1e9 if pd.isna(r.upper_limit) else r.upper_limit])
nS = len(s) + 1
wm = wb.create_sheet("measurement")
wm.append(["unit_id", "test_date", "model", "parameter", "value", "valid_from", "lower", "upper", "pass"])
for i, r in enumerate(m.itertuples(), start=2):
    wm.append([r.unit_id, r.test_date.to_pydatetime(), r.model, r.parameter, r.value,
               f'=_xlfn.MAXIFS(spec!$C$2:$C${nS},spec!$A$2:$A${nS},C{i},spec!$B$2:$B${nS},D{i},spec!$C$2:$C${nS},"<="&B{i})',
               f"=SUMIFS(spec!$E$2:$E${nS},spec!$A$2:$A${nS},C{i},spec!$B$2:$B${nS},D{i},spec!$C$2:$C${nS},F{i})",
               f"=SUMIFS(spec!$F$2:$F${nS},spec!$A$2:$A${nS},C{i},spec!$B$2:$B${nS},D{i},spec!$C$2:$C${nS},F{i})",
               f"=IF(AND(E{i}>=G{i},E{i}<=H{i}),1,0)"])
nM = len(m) + 1
wc = wb.create_sheet("check", 0)
wc.append(["Cek silang: sel di bawah adalah RUMUS Excel dari lembar measurement dan spec"]); wc["A1"].font = Font(bold=True)
wc.append(["parameter", "model", "n", "n_pass", "pass_rate"])
for r in ref.itertuples():
    k = wc.max_row + 1
    wc.append([r.parameter, r.model, f'=COUNTIFS(measurement!$D$2:$D${nM},A{k},measurement!$C$2:$C${nM},B{k})',
               f'=SUMIFS(measurement!$I$2:$I${nM},measurement!$D$2:$D${nM},A{k},measurement!$C$2:$C${nM},B{k})', f"=D{k}/C{k}"])
for sh in (ws, wm, wc):
    for c in sh[1 if sh is not wc else 2]:
        c.font = Font(bold=True)
tmp = Path(tempfile.mkdtemp()); src = tmp / "in" / "crosscheck.xlsx"; src.parent.mkdir()
wb.save(src)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(src)], check=True, capture_output=True, timeout=580)
shutil.copy(tmp / "out" / "crosscheck.xlsx", C.RES_ANALYSIS / "crosscheck.xlsx")
v = load_workbook(C.RES_ANALYSIS / "crosscheck.xlsx", data_only=True)["check"]
rows = []
for r, p in zip(v.iter_rows(min_row=3, max_row=2 + len(ref), values_only=True), ref.itertuples()):
    for name, xl, py, tol in (("n", r[2], p.n, 0), ("n_pass", r[3], p.n_pass, 0), ("pass_rate", r[4], p.pass_rate, 1e-12)):
        rows.append((p.parameter, p.model, name, xl, py, "PASS" if xl is not None and abs(xl - py) <= tol else "FAIL"))
df = pd.DataFrame(rows, columns=["parameter", "model", "metric", "excel", "sql", "status"])
df.to_csv(C.RES_ANALYSIS / "crosscheck_excel.csv", index=False)
print(df.status.value_counts().to_dict())
assert (df.status == "PASS").all(), df[df.status == "FAIL"]
